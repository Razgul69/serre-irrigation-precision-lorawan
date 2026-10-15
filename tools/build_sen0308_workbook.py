"""Build the SEN0308 calibration workbook from the CR1000 TOA5 export."""
from __future__ import annotations

import csv
import datetime as dt
import math
import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


REPO = Path(__file__).resolve().parents[1]
SOURCE = Path(r"D:\Mes Programmes\CR1000calibsonde_CalibSEN0308.dat")
REPO_OUTPUT = REPO / "docs" / "Calibration_SEN0308.xlsx"
LOCAL_OUTPUT = Path(r"D:\capteur arduino\Calibration_SEN0308.xlsx")
DRY_POT_SOIL_G = 2050.0
VOLUME_ML = 4800.0
FIELD_CAPACITY_G = 3555.0
REPEAT_WEIGHT_G = 3530.0
CURRENT_WEIGHT_G = 3523.0
CURRENT_WINDOW_END = dt.datetime(2026, 10, 15, 14, 44, 30)

NAVY = "17324D"
TEAL = "147D82"
PALE = "EAF2F4"
PALE_YELLOW = "FFF2CC"
WHITE = "FFFFFF"
GRAY = "607080"


def read_toa5(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.reader(stream))
    headers = rows[1]
    result = []
    for values in rows[4:]:
        if not values or not values[0].startswith("20"):
            continue
        result.append(dict(zip(headers, values)))
    return result


def summarize_window(rows: list[dict[str, str]], start: dt.datetime, end: dt.datetime) -> dict:
    selected = [
        row for row in rows
        if start < dt.datetime.strptime(row["TIMESTAMP"], "%Y-%m-%d %H:%M:%S") <= end
    ]
    values = [float(row["SEN_mV_Avg"]) for row in selected]
    block_sds = [float(row["SEN_mV_Std"]) for row in selected]
    if not values:
        raise ValueError(f"No CR1000 records in {start} to {end}")
    mean = sum(values) / len(values)
    variance = sum(sd * sd + (value - mean) ** 2 for value, sd in zip(values, block_sds)) / len(values)
    return {
        "start": start,
        "end": end,
        "count": len(values),
        "mean": mean,
        "sd": math.sqrt(variance),
        "max_block_sd": max(block_sds),
    }


def style_title(ws, title: str, subtitle: str, last_col: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(1, 1, title)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name="Aptos Display", size=18, bold=True, color=WHITE)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    cell = ws.cell(2, 1, subtitle)
    cell.font = Font(name="Aptos", size=10, italic=True, color=GRAY)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 30


def style_header(ws, row: int, last_col: int) -> None:
    for col in range(1, last_col + 1):
        cell = ws.cell(row, col)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(name="Aptos", bold=True, color=WHITE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 32


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(f"CR1000 export not found: {SOURCE}")
    rows = read_toa5(SOURCE)
    point0 = [row for row in rows if row["Point"] == "0"]
    point6 = [row for row in rows if row["Point"] == "6"]
    dry_window = summarize_window(
        point0, dt.datetime(2026, 10, 12, 17, 35, 30), dt.datetime(2026, 10, 12, 17, 40, 30)
    )
    fc_window = summarize_window(
        point6, dt.datetime(2026, 10, 14, 9, 19), dt.datetime(2026, 10, 14, 9, 24)
    )
    latest_window = summarize_window(
        point6, CURRENT_WINDOW_END - dt.timedelta(minutes=5), CURRENT_WINDOW_END
    )
    final_record = point6[-1]

    wb = Workbook()
    readme = wb.active
    readme.title = "Lisez-moi"
    points = wb.create_sheet("Points calibration")
    raw = wb.create_sheet("Données CR1000")

    style_title(readme, "Calibration de la sonde DFRobot SEN0308", "Carnet de pré-manipulation — terreau 4,8 L, CR1000, export mis à jour le 15/10/2026", 4)
    readme.sheet_view.showGridLines = False
    for col, width in {"A": 29, "B": 34, "C": 34, "D": 34}.items():
        readme.column_dimensions[col].width = width
    readme_rows = [
        ("Objet", "Rassembler les masses, les relevés SEN0308/CR1000 et les notes de calibration. Les données brutes exportées sont dans l’onglet « Données CR1000 »."),
        ("Montage", "SEN0308 : rouge → 5 V CR1000, noir → G, jaune → entrée 1H (SE1). Ne jamais mettre la sonde à l’étuve."),
        ("Pot", "Pot vide 140 g ; volume nominal 4 800 mL."),
        ("Référence sèche corrigée", "Pot + terreau sec étuvé à 110 °C : 2 050 g, soit 1 910 g de terreau sec. L’ancienne valeur 2 256,57 g a été corrigée et n’est pas utilisée."),
        ("Point capacité au champ", "Pot drainé après arrosage : 3 555 g ; eau retenue estimée 1 505 g ; θ apparente = 31,35 % vol. Point mesuré le 14/10, après drainage."),
        ("Pesée précédente", "Pot pesé à 3 530 g le 15/10 : 1 480 g au-dessus de la référence sèche, soit 30,83 % vol. Pas de moyenne mV exactement associée à cette pesée."),
        ("Formule gravimétrique", "θ (% vol.) = 100 × (masse pot+terreau humide − 2 050 g) / 4 800 mL. Approximation basée sur 1 g d’eau ≈ 1 mL et le volume nominal du pot."),
        ("Point sec — lecture", f"Fenêtre 12/10 17:35:30–17:40:30 : {dry_window['mean']:.2f} mV, {dry_window['count']} blocs de 30 s, écart-type global estimé {dry_window['sd']:.2f} mV."),
        ("Capacité au champ — lecture", f"Fenêtre 14/10 09:19–09:24 : {fc_window['mean']:.2f} mV, {fc_window['count']} blocs, écart-type global estimé {fc_window['sd']:.2f} mV. C’est le relevé associé à la capacité au champ."),
        ("Mesure actuelle — provisoire", f"Poids déclaré : {CURRENT_WEIGHT_G:.0f} g. Fenêtre stable {latest_window['start']:%d/%m %H:%M:%S}–{latest_window['end']:%H:%M:%S} : {latest_window['mean']:.2f} mV (10 blocs, écart-type estimé {latest_window['sd']:.2f} mV). θ apparente = {100 * (CURRENT_WEIGHT_G - DRY_POT_SOIL_G) / VOLUME_ML:.2f} % vol."),
        ("Contrôle de cohérence", f"La masse actuelle est seulement {FIELD_CAPACITY_G-CURRENT_WEIGHT_G:.0f} g sous la capacité au champ, mais le signal diffère de {latest_window['mean']-fc_window['mean']:.1f} mV. Point à vérifier (même profondeur/contact, pot et câble pesés de la même façon) ; ne pas encore l’utiliser dans l’ajustement."),
        ("Bloc final rejeté", f"Le dernier enregistrement du fichier est {final_record['TIMESTAMP']} : {float(final_record['SEN_mV_Avg']):.2f} mV, écart-type {float(final_record['SEN_mV_Std']):.2f} mV. Il est exclu de la moyenne stable."),
        ("État de calibration", "Deux références restent utilisables (terreau sec et capacité au champ). Le nouveau couple poids/mV est conservé comme provisoire, car sa variation de signal paraît incohérente avec l’écart de masse."),
        ("Prochain point utile", "Laisser sécher sans ajouter d’eau. Pour une baisse d’environ 10 % vol. depuis la capacité au champ : viser ≈ 3 075 g pot+terreau (sans sonde), puis noter la masse et la moyenne mV sur une fenêtre stable de 5 min."),
        ("Comment ajouter un relevé", "Dans « Points calibration », saisir une ligne avec le poids du pot sans sonde, la date/heure, la moyenne SEN_mV sur 5 min et le plus grand écart-type des blocs. Les colonnes d’eau retenue et θ sont calculées automatiquement."),
        ("Attention au champ Point", "Le programme a gardé Point=6 du 14/10 au 15/10. Ce numéro ne prouve pas que toutes ces lignes sont au même état d’humidité : utiliser poids + heure pour identifier les étapes. Le bloc de 14:45 est exclu."),
        ("Source", str(SOURCE)),
    ]
    for index, (label, text) in enumerate(readme_rows, start=4):
        readme.cell(index, 1, label).font = Font(bold=True, color=NAVY)
        readme.cell(index, 1).fill = PatternFill("solid", fgColor=PALE)
        readme.cell(index, 2, text).alignment = Alignment(wrap_text=True, vertical="top")
        readme.merge_cells(start_row=index, start_column=2, end_row=index, end_column=4)
        readme.row_dimensions[index].height = 42 if len(text) > 140 else 30
    readme.freeze_panes = "A4"

    point_headers = ["Point / état", "Date / fenêtre de 5 min", "Poids pot + terreau (g), sans sonde", "Masse sèche de référence (g)", "Eau retenue (g)", "θ apparente (% vol.)", "SEN0308 moyenne (mV)", "Écart-type fenêtre estimé (mV)", "Blocs 30 s", "Statut / note"]
    style_title(points, "Points de calibration et relevés manuels", "Les fenêtres déjà calculées sont séparées des pesées sans mV confirmé. Compléter les lignes jaunes au fil du dessèchement.", len(point_headers))
    for col, header in enumerate(point_headers, 1):
        points.cell(4, col, header)
    style_header(points, 4, len(point_headers))
    fixed = [
        ["0 — terreau sec étuvé", "12/10/2026 17:35:30–17:40:30", 2050.0, DRY_POT_SOIL_G, None, None, dry_window["mean"], dry_window["sd"], dry_window["count"], "Mesure brute stable ; masse sèche corrigée le 15/10."],
        ["6 — capacité au champ drainée", "14/10/2026 09:19–09:24", FIELD_CAPACITY_G, DRY_POT_SOIL_G, None, None, fc_window["mean"], fc_window["sd"], fc_window["count"], "Poids déclaré 3 555 g ; lecture stable associée."],
        ["Répétition / séchage", "15/10/2026 — heure non précisée", REPEAT_WEIGHT_G, DRY_POT_SOIL_G, None, None, None, None, None, "Pesée déclarée ; pas de moyenne mV associée avec certitude."],
        ["6 — séchage, à vérifier", f"15/10/2026 {latest_window['start']:%H:%M:%S}–{latest_window['end']:%H:%M:%S}", CURRENT_WEIGHT_G, DRY_POT_SOIL_G, None, None, latest_window["mean"], latest_window["sd"], latest_window["count"], f"Couple déclaré 3 523 g / {latest_window['mean']:.2f} mV ; provisoire : +{latest_window['mean']-fc_window['mean']:.1f} mV pour seulement -{FIELD_CAPACITY_G-CURRENT_WEIGHT_G:.0f} g vs capacité au champ."],
        ["Bloc isolé, rejeté", final_record["TIMESTAMP"], None, DRY_POT_SOIL_G, None, None, float(final_record["SEN_mV_Avg"]), float(final_record["SEN_mV_Std"]), 1, "Écart-type élevé au dernier enregistrement ; ne pas utiliser."],
        ["5 — prochain palier visé", "À compléter", 3075.0, DRY_POT_SOIL_G, None, None, None, None, None, "Cible indicative, environ 10 % vol. sous capacité au champ ; saisir la vraie pesée et la fenêtre stable."],
        ["4 — palier suivant indicatif", "À compléter", 2595.0, DRY_POT_SOIL_G, None, None, None, None, None, "Cible indicative ; saisir la vraie pesée et la fenêtre stable."],
    ]
    for r, values in enumerate(fixed, start=5):
        for c, value in enumerate(values, start=1):
            if value is not None:
                points.cell(r, c, value)
        points.cell(r, 5, f'=IF(C{r}="","",C{r}-D{r})')
        points.cell(r, 6, f'=IF(E{r}="","",100*E{r}/{VOLUME_ML})')
        if r in (7, 8, 10, 11):
            for cell in points[r]:
                cell.fill = PatternFill("solid", fgColor=PALE_YELLOW)
    for _ in range(6):
        r = points.max_row + 1
        points.cell(r, 1, "À compléter")
        points.cell(r, 4, DRY_POT_SOIL_G)
        points.cell(r, 5, f'=IF(C{r}="","",C{r}-D{r})')
        points.cell(r, 6, f'=IF(E{r}="","",100*E{r}/{VOLUME_ML})')
        for cell in points[r]:
            cell.fill = PatternFill("solid", fgColor=PALE_YELLOW)
    for row in range(5, points.max_row + 1):
        for col in (3, 4, 5):
            points.cell(row, col).number_format = '#,##0.00'
        points.cell(row, 6).number_format = '0.00'
        for col in (7, 8):
            points.cell(row, col).number_format = '0.000'
    widths = [28, 30, 25, 23, 18, 21, 22, 27, 13, 60]
    for col, width in enumerate(widths, 1):
        points.column_dimensions[chr(64 + col)].width = width
    points.freeze_panes = "A5"
    points.auto_filter.ref = f"A4:J{points.max_row}"
    points.sheet_view.showGridLines = False

    raw_headers = ["TIMESTAMP", "RECORD", "SEN_mV_Avg", "SEN_mV_Std", "SEN_mV_Min", "SEN_mV_Max", "Point logger", "PTemp_C"]
    style_title(raw, "Export brut CR1000 — CalibSEN0308", f"{len(rows)} lignes importées depuis {SOURCE.name}. Valeurs inchangées, filtres disponibles.", len(raw_headers))
    for col, header in enumerate(raw_headers, 1):
        raw.cell(4, col, header)
    style_header(raw, 4, len(raw_headers))
    for r, item in enumerate(rows, start=5):
        raw.cell(r, 1, dt.datetime.strptime(item["TIMESTAMP"], "%Y-%m-%d %H:%M:%S"))
        raw.cell(r, 1).number_format = "yyyy-mm-dd hh:mm:ss"
        raw.cell(r, 2, int(item["RECORD"]))
        for col, key in enumerate(raw_headers[2:6], start=3):
            raw.cell(r, col, float(item[key]))
            raw.cell(r, col).number_format = "0.000"
        raw.cell(r, 7, int(item["Point"]))
        raw.cell(r, 8, float(item["PTemp_C"]))
        raw.cell(r, 8).number_format = "0.00"
    last_raw = 4 + len(rows)
    table = Table(displayName="CR1000Data", ref=f"A4:H{last_raw}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showColumnStripes=False)
    raw.add_table(table)
    raw.freeze_panes = "A5"
    raw.sheet_view.showGridLines = False
    for col, width in enumerate([22, 12, 16, 16, 16, 16, 15, 14], 1):
        raw.column_dimensions[chr(64 + col)].width = width

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.active = 0
    REPO_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(REPO_OUTPUT)
    LOCAL_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REPO_OUTPUT, LOCAL_OUTPUT)

    check = load_workbook(REPO_OUTPUT, read_only=True, data_only=False)
    assert check.sheetnames == ["Lisez-moi", "Points calibration", "Données CR1000"]
    assert check["Données CR1000"].max_row == last_raw
    check.close()
    print(f"Created {REPO_OUTPUT} and {LOCAL_OUTPUT}; {len(rows)} raw records; "
          f"dry={dry_window['mean']:.3f} mV, field capacity={fc_window['mean']:.3f} mV, "
          f"latest unvalidated={latest_window['mean']:.3f} mV")


if __name__ == "__main__":
    main()