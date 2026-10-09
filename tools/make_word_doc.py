# -*- coding: utf-8 -*-
"""Genere le document Word du projet Irrigation de Precision LoRaWAN."""
from docx import Document
from docx.shared import Pt, RGBColor, Cm

doc = Document()

def h(text, level=1):
    doc.add_heading(text, level=level)

def p(text, bold=False):
    par = doc.add_paragraph()
    run = par.add_run(text)
    run.bold = bold
    return par

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, htxt in enumerate(headers):
        t.rows[0].cells[i].text = htxt
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    doc.add_paragraph()

# Titre
title = doc.add_heading("Irrigation de Précision LoRaWAN — Serre", level=0)
p("Système d'irrigation piloté par le potentiel matriciel du sol, inspiré du projet "
  "Makerfabs IoT Irrigation System (github.com/Makerfabs/Project_IoT-Irrigation-System).")
p("Dépôt GitHub : https://github.com/Razgul69/serre-irrigation-precision-lorawan", bold=True)

h("1. Architecture générale")
p("Chaîne de mesure : sondes TEROS 21 (METER) → convertisseurs Dragino SDI-12 → LoRaWAN → "
  "passerelle LoRa dans la serre → The Things Network (TTN) → serveur virtuel.")
p("Chaîne de commande : Node-RED (orchestrateur) évalue les seuils d'irrigation et envoie des "
  "downlinks LoRaWAN vers des ESP32 qui pilotent une électrovanne par sonde (irrigation zonale). "
  "Un débitmètre à impulsions sur chaque ESP32 mesure l'apport d'eau exact.")
p("Stockage et visualisation : InfluxDB + Grafana (tensiométrie, volumes, états des vannes, alertes).")

h("2. Matériel")
table(["Élément", "Référence", "Rôle"], [
    ["Sonde tensiométrique", "METER TEROS 21", "Potentiel matriciel + température sol (SDI-12)"],
    ["Sonde capacitive (complément)", "DFRobot Gravity SEN0308 V2.0", "Indice d'humidité relatif / teneur en eau après calibration (0-3 V) ; sur CR1000 (1H) ou Nano ESP32 (A0)"],
    ["Convertisseur", "Dragino SDI-12-LB", "SDI-12 vers LoRaWAN, alimenté batterie"],
    ["Passerelle", "Passerelle LoRaWAN 868 MHz", "Uplink/downlink vers TTN"],
    ["Contrôleur vanne", "ESP32 + RFM95/SX1276", "Pilotage électrovanne + comptage débit"],
    ["Électrovanne", "12/24 V (1 par sonde/zone)", "Irrigation zonale de précision"],
    ["Débitmètre", "À impulsions (ex. YF-S201)", "Mesure de l'apport d'eau"],
    ["Serveur", "VPS Docker", "Node-RED + InfluxDB + Grafana"],
])

h("3. Logique d'irrigation")
table(["Paramètre", "Valeur par défaut", "Description"], [
    ["SEUIL_SEC", "-40 kPa", "En dessous : démarrer l'irrigation"],
    ["SEUIL_HUMIDE", "-10 kPa", "Au-dessus : arrêter l'irrigation"],
    ["VOLUME_MAX", "10 L", "Sécurité : volume max par cycle"],
    ["DUREE_MAX", "15 min", "Sécurité : durée max d'ouverture"],
])
p("Sécurités embarquées dans l'ESP32 (autonomes même en cas de perte réseau) : "
  "fermeture automatique sur durée max ou volume max atteint.")
p("Garde-fous côté serveur : anti-rebond 30 min entre cycles, alerte si vanne ouverte sans débit "
  "(vanne HS/fuite), alerte stress hydrique si potentiel < -80 kPa.")

h("3 bis. Pré-manips — banc de laboratoire (avant déploiement LoRaWAN)")
p("Avant l'installation en serre, les capteurs sont caractérisés sur un banc filaire : un pot sur balance, "
  "une TEROS 21/22 sur datalogger Campbell, et la sonde capacitive DFRobot SEN0308 en comparaison. "
  "Objectif : relation poids (évapotranspiration) / potentiel matriciel / indice SEN0308 par substrat, "
  "pour fixer les seuils d'irrigation. Ces éléments ne sont pas requis pour le système final.")
p("Balance Ohaus DP15/M — interfaces PC et Mac (dossier balance/)", bold=True)
table(["Plateforme", "Fichier", "Fonctionnement"], [
    ["Windows (PC)", "balance_gui.py → BalanceCollecteur.exe", "Fenêtre Marche/Arrêt, détection auto du port USB, lecture SICS toutes les 5 s, CSV incrémental, envoi InfluxDB local, veille PC bloquée"],
    ["macOS", "balance_gui_macos.py, install_mac.sh, build_macos.sh", "Même interface, CSV local uniquement, indicateurs port détecté / port ouvert / balance connectée ; install_mac.sh installe Python 3.12 + Tk 8.6"],
    ["Ligne de commande", "collect_balance.py", "python collect_balance.py --port COM3 --interval 5 [--no-influx]"],
    ["Diagnostic", "diagnose_serial.py, diagnose_serial_macos.py, DIAGNOSTIC.md", "Test adaptateur / port / réponse balance étape par étape"],
])
p("Datalogger Campbell (dossier campbell/)", bold=True)
table(["Programme", "Cible", "Contenu"], [
    ["teros22_cr1000.CR1", "CR1000", "TEROS 22 seule (SDI-12 sur C3), table 15 min"],
    ["teros22_sen0308_cr1000.CR1", "CR1000", "TEROS 22 + SEN0308 (5V / G / 1H), tables MesuresSol + CalibSEN0308"],
    ["teros22_sen0308_cr800.CR8", "CR800", "Identique pour CR800"],
])
p("La balance est volontairement hors datalogger : acquise par le PC/Mac et croisée par horodatage.")

h("3 ter. Pré-manip — Sonde DFRobot SEN0308 : calibration et intégration")
p("La SEN0308 est une sonde capacitive analogique (0-3 V, décroissante avec l'humidité). Elle fournit un "
  "indice relatif dépendant du substrat, du tassement et de la tension d'alimentation : elle doit être "
  "calibrée sur la chaîne de mesure définitive. Protocole complet : docs/calibration_sen0308.md.")
table(["Niveau", "Résultat", "Méthode"], [
    ["1 — Air / eau", "Indice 0-100 %", "Moyenne 30 s sonde sèche dans l'air (AIR_MV) puis dans l'eau à la ligne d'insertion (WATER_MV)"],
    ["2 — Gravimétrique", "Teneur en eau volumique θ (% vol)", "5-7 pesées (balance Ohaus) d'un pot de volume connu du saturé au sec étuve ; droite θ = a·mV + b (R² ≥ 0,95)"],
    ["3 — Co-localisation TEROS", "Seuils d'irrigation sur l'indice", "SEN0308 et TEROS côte à côte sur 1-2 cycles ; lire l'indice à -40 et -10 kPa"],
])
table(["Option", "Câblage", "Fichiers"], [
    ["A — Campbell CR1000 / CR800 (recommandé)", "rouge→5V, noir→G, jaune→1H", "campbell/teros22_sen0308_cr1000.CR1, campbell/teros22_sen0308_cr800.CR8"],
    ["B — Arduino Nano ESP32", "rouge→3V3, noir→GND, jaune→A0", "firmware/esp32_nano_sen0308/, tools/calib_sen0308.py, ttn/decoder_sen0308.js (variante LoRaWAN, port 3)"],
])

h("4. Protocole LoRaWAN")
p("Uplink contrôleur vanne (port 2) : état vanne (1 o), volume du cycle en mL (4 o), "
  "volume total en L x10 (2 o), uptime en minutes (2 o).")
p("Downlink (port 1) : 0x00 = fermer ; 0x01 = ouvrir ; 0x02 + volume/100 mL (2 o) = "
  "ouvrir avec limite de volume.")
p("Uplink Dragino/TEROS21 : batterie (mV) + réponse SDI-12 ASCII "
  "(potentiel matriciel kPa + température °C), décodée par le payload formatter TTN.")
p("Uplink Nano ESP32/SEN0308 (port 3) : tension brute mV (2 o), indice x10 (2 o), θ x10 (2 o, 0xFFFF si non calibré).")

h("5. Structure du dépôt")
table(["Dossier", "Contenu"], [
    ["firmware/esp32_valve_controller/", "Firmware PlatformIO (LMIC, vanne, débitmètre, sécurités)"],
    ["firmware/esp32_nano_sen0308/", "Firmware Nano ESP32 : SEN0308, CSV série, capture air/eau, option LoRaWAN"],
    ["balance/", "[Pré-manip] Collecteur balance Ohaus : interfaces Windows et macOS, CLI, diagnostics"],
    ["campbell/", "[Pré-manip] Programmes CR1000 / CR800 : TEROS 22 seule ou TEROS 22 + SEN0308"],
    ["tools/calib_sen0308.py", "Enregistrement série SEN0308 (+InfluxDB) et ajustement gravimétrique"],
    ["docs/calibration_sen0308.md", "Protocole de calibration SEN0308 (3 niveaux) et fiche de relevé"],
    ["ttn/", "Décodeurs de payload TTN (TEROS21, contrôleur vanne, SEN0308)"],
    ["server/docker-compose.yml", "Stack Node-RED + InfluxDB 2.x + Grafana"],
    ["server/nodered/", "Documentation du flow d'orchestration et des seuils par zone"],
    ["grafana/", "Panels recommandés et requêtes Flux"],
    ["docs/images/", "Photos du montage (serre, sondes, vannes, passerelle)"],
])

h("6. Mise en route")
for step in [
    "1. TTN : créer une application, enregistrer les Dragino et les ESP32 en OTAA, coller les décodeurs.",
    "2. Serveur : docker compose up -d, configurer l'intégration MQTT TTN dans Node-RED.",
    "3. Firmware : renseigner les clés OTAA dans config.h, compiler avec PlatformIO, flasher.",
    "4. Grafana : créer la datasource InfluxDB et les panels (voir grafana/README.md).",
    "5. Calibrer les seuils par zone selon la culture et le substrat.",
]:
    p(step)

h("7. État d'avancement (09/10/2026)")
table(["Étape", "État"], [
    ["Pré-manip — collecteur balance Ohaus (Windows exe + macOS), CSV + InfluxDB, dashboard Grafana poids", "OPÉRATIONNEL (PC) ; Mac validé en diagnostic série"],
    ["Pré-manip — programmes Campbell CR1000/CR800 TEROS 22 (+ SEN0308)", "ÉCRITS (TEROS seule testée sur CR800)"],
    ["Pré-manip — sonde SEN0308 : protocole, firmware Nano ESP32, outil Python", "ÉCRITS — calibration terrain À FAIRE"],
    ["Firmware ESP32 vanne + débitmètre", "FAIT (à flasher/tester)"],
    ["Décodeurs TTN (TEROS21, vanne, SEN0308)", "FAIT"],
    ["Stack Docker Node-RED/InfluxDB/Grafana", "OPÉRATIONNELLE en local (ports 11880/18086/13000)"],
    ["Flux Node-RED « IoT Irrigation Piloté »", "DÉPLOYÉ : dashboard /ui avec jauge kPa, seuils réglables, mode auto, boutons OUVRIR/FERMER par zone, écriture InfluxDB, downlinks TTN"],
    ["Export du flux (importable)", "server/nodered/flow_iot_irrigation_pilote.json"],
    ["Application TTN", "À CRÉER (identifiants MQTT à renseigner dans Node-RED)"],
    ["Enregistrement devices (Dragino, ESP32)", "À FAIRE"],
    ["Installation terrain (sondes, vannes, passerelle)", "À FAIRE"],
    ["Photos du montage", "À AJOUTER dans docs/images/"],
])

doc.save(r"D:\Projet_Irrigation_LoRaWAN\Projet_Irrigation_Precision_LoRaWAN.docx")
print("OK: D:\\Projet_Irrigation_LoRaWAN\\Projet_Irrigation_Precision_LoRaWAN.docx")
